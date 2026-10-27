"""文档解析：把各种格式统一成「带 # 标题的 Markdown 风格文本」。

统一格式的意义：分块器只需要理解一种输入，标题层级就是分块边界，
这样「第三章 考勤 > 3.2 请假流程」这种结构才能保留到引用里。

两条抽取路线：
- **文本层**：md / txt / csv / html / docx / pptx / xlsx / 有文本层的 pdf ——直接读字符；
- **识别类**：图片、扫描版 pdf ——先 OCR 再读字符。

识别类内容必须留痕（`meta.ocr` / `meta.ocr_confidence`）：
它的文字是"猜"出来的，检索时要降权，而且识别不出字时不能返回空字符串假装成功。
"""
from __future__ import annotations

import html as html_lib
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Tuple

log = logging.getLogger(__name__)

TEXT_EXTS = {".md", ".markdown", ".txt", ".csv", ".tsv"}
OFFICE_EXTS = {".docx", ".pptx", ".xlsx", ".xlsm"}
PDF_EXTS = {".pdf"}
HTML_EXTS = {".html", ".htm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
SUPPORTED_EXTS = TEXT_EXTS | OFFICE_EXTS | PDF_EXTS | HTML_EXTS | IMAGE_EXTS

# 明确"认识但暂不支持"的格式——必须能报出来，不能静默跳过
AUDIO_VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".flv", ".wmv", ".m4v", ".webm",
                    ".mp3", ".wav", ".m4a", ".aac", ".flac", ".wma", ".ogg", ".amr"}
LEGACY_OFFICE_EXTS = {".doc", ".xls", ".ppt", ".dot", ".wps", ".et", ".dps"}
ARCHIVE_EXTS = {".zip", ".rar", ".7z", ".tar", ".gz", ".bz2"}

PAGE_BREAK = "\f"
_HEADER_BAND = 0.08     # 页面上 8% 视为页眉区
_FOOTER_BAND = 0.92     # 页面下 8% 视为页脚区
_MIN_TEXT_PER_PAGE = 20  # 少于此字符数视为「没有文本层」，走 OCR


@dataclass
class LoadedDoc:
    text: str
    title: str
    source_type: str
    meta: Dict[str, object] = field(default_factory=dict)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 零宽字符与软连字符：从 PDF / 网页复制粘贴时带进来的，肉眼看不见但会切断词
    text = re.sub(r"[\u200b\u200c\u200d\ufeff\u00ad]", "", text)
    text = re.sub(r"[ \t\u00a0\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_file(path: Path) -> LoadedDoc:
    ext = path.suffix.lower()
    if ext in (".md", ".markdown"):
        return _load_markdown(path)
    if ext in (".txt", ".csv", ".tsv"):
        return _load_text(path)
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".pptx":
        return _load_pptx(path)
    if ext in (".xlsx", ".xlsm"):
        return _load_xlsx(path)
    if ext in HTML_EXTS:
        return _load_html(path)
    if ext in IMAGE_EXTS:
        return _load_image(path)
    raise ValueError(f"不支持的格式: {ext}")


# ---------------- 文本类 ----------------

# 注意：分隔行可能是 CRLF（Windows 上生成的文件），\r 不在 [ \t] 里，
# 必须显式容忍 \r?\n，否则 front-matter 会被当成正文原样入库。
_FM_RE = re.compile(r"\A---[ \t]*\r?\n(.*?)\r?\n---[ \t]*(?:\r?\n|\Z)", re.S)


def _parse_front_matter(text: str) -> Tuple[Dict[str, object], str]:
    """剥掉 md 头部的 YAML front-matter，返回 (元数据, 正文)。

    只覆盖我们实际会写的简单形态（key: value 单行），
    不引第三方 YAML 依赖——这个模块要保持"任何环境都能跑"。
    """
    m = _FM_RE.match(text)
    if not m:
        return {}, text
    meta: Dict[str, object] = {}
    for line in m.group(1).splitlines():
        s = line.strip()
        if not s or s.startswith("#") or ":" not in s:
            continue
        k, v = s.split(":", 1)
        k, v = k.strip(), v.strip()
        if len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
            v = v[1:-1]
        if v.lower() in ("true", "false"):
            meta[k] = v.lower() == "true"
        elif re.fullmatch(r"-?\d+", v):
            meta[k] = int(v)
        else:
            meta[k] = v
    return meta, text[m.end():]


def _load_markdown(path: Path) -> LoadedDoc:
    raw = _read_text(path)
    meta, text = _parse_front_matter(raw)
    title = str(meta.get("title") or "").strip() or path.stem
    m = re.search(r"^#\s+(.+?)\s*$", text, re.M)
    if m:
        title = m.group(1).strip()
    meta.setdefault("modality", "text")
    return LoadedDoc(_clean(text), title, "markdown", meta)


def _load_text(path: Path) -> LoadedDoc:
    text = _read_text(path)
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv"):
        sep = "\t" if ext == ".tsv" else ","
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        rows = ["| " + " | ".join(c.strip() for c in ln.split(sep)) + " |" for ln in lines]
        return LoadedDoc(_clean(f"## {path.stem}\n" + "\n".join(rows)), path.stem,
                        "csv", {"modality": "text"})
    return LoadedDoc(_clean(_infer_headings_from_plain(text)), path.stem, "text",
                     {"modality": "text"})


def _infer_headings_from_plain(text: str) -> str:
    """纯文本里识别「第一章 / 第 1 章 / 一、 / 1.1」这类标题行，补上 # 让分块器能用。"""
    out: List[str] = []
    pat = re.compile(
        r"^(第\s*[一二三四五六七八九十百0-9]+\s*[章节部分篇]\s*.+|"
        r"[一二三四五六七八九十]+、.{2,40}|"
        r"\d+(?:\.\d+){0,2}\s+\S.{0,40})$")
    for line in text.split("\n"):
        s = line.strip()
        if s and len(s) <= 45 and pat.match(s) and not s.endswith("。"):
            out.append(f"## {s}")
        else:
            out.append(line)
    return "\n".join(out)


# ---------------- PDF ----------------

@dataclass
class _Page:
    index: int
    width: float
    height: float
    blocks: List[tuple]
    ocr_text: str = ""
    ocr_conf: float = 0.0


def _load_pdf(path: Path) -> LoadedDoc:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        log.warning("未安装 pymupdf，PDF 退回 pypdf：版面顺序、页眉页脚、扫描页都无法处理")
        return _load_pdf_pypdf(path)

    doc = fitz.open(str(path))
    try:
        title = path.stem
        try:
            if doc.metadata and doc.metadata.get("title"):
                title = str(doc.metadata["title"]).strip() or title
        except Exception:  # noqa: BLE001
            pass

        pages: List[_Page] = []
        for i, page in enumerate(doc):
            blocks = [b for b in page.get_text("blocks")
                      if len(b) > 6 and b[6] == 0 and str(b[4]).strip()]
            pages.append(_Page(i, page.rect.width, page.rect.height,
                               sorted(blocks, key=lambda b: (b[1], b[0]))))

        _drop_running_heads(pages)

        ocr_pages = 0          # 真正 OCR 出文字的页数
        ocr_attempts = 0       # 尝试 OCR 的页数（含失败）——两者必须分开记，
                               # 否则"OCR 跑了但一个字没认出来"会被当成"没跑过"
        confs: List[float] = []
        for p in pages:
            if _text_len(p.blocks) >= _MIN_TEXT_PER_PAGE:
                continue
            ocr_attempts += 1
            text, conf, n = _ocr_page(doc[p.index], path.name, p.index)
            if n:
                p.ocr_text, p.ocr_conf = text, conf
                ocr_pages += 1
                confs.append(conf)

        parts: List[str] = []
        for p in pages:
            if p.ocr_text:
                parts.append(p.ocr_text)
            else:
                parts.append(_blocks_to_text(p))

        text = PAGE_BREAK.join(parts)
        text = re.sub(r"(?<=[。！？；])\s*(?=[^\n])", "\n", text)
        meta: Dict[str, object] = {
            "modality": "pdf",
            "pages": len(pages),
            "ocr_pages": ocr_pages,
            "ocr_attempted_pages": ocr_attempts,
            "ocr": ocr_attempts > 0,
        }
        if confs:
            meta["ocr_confidence"] = round(sum(confs) / len(confs), 4)
        if not text.replace(PAGE_BREAK, "").strip():
            if ocr_attempts:
                meta["empty_reason"] = (
                    f"扫描版 PDF，{len(pages)} 页都未识别出文字"
                    "（可能分辨率过低、图片本身无字，或整页都是图形）")
            else:
                meta["empty_reason"] = "PDF 未抽出任何文字（可能是纯图形页或加密文档）"
        return LoadedDoc(_clean(_infer_headings_from_plain(text)), title, "pdf", meta)
    finally:
        doc.close()


def _text_len(blocks: List[tuple]) -> int:
    return sum(len(str(b[4]).strip()) for b in blocks)


def _drop_running_heads(pages: List[_Page]) -> None:
    """删掉在多页重复出现的页眉/页脚。

    判定方式：位于页面上下 8% 的内容，去掉数字后若在多页重复出现，就是页眉页脚。
    页码单独处理：纯粹的「- 3 -」「第 3 页」一律删掉——页码的价值是定位，
    而定位已经由 PAGE_BREAK 记在 meta.page 里了，正文里再留一份只是噪声。
    """
    if not pages:
        return
    counts: Dict[str, int] = {}
    for p in pages:
        for key in {_band_key(b[4]) for b in p.blocks if _in_band(b, p)}:
            if key:
                counts[key] = counts.get(key, 0) + 1
    threshold = max(3, int(len(pages) * 0.5))
    repeated = {k for k, c in counts.items() if c >= threshold}

    for p in pages:
        kept = []
        for b in p.blocks:
            txt = str(b[4]).strip()
            if _is_page_number(txt):
                continue
            if _in_band(b, p) and _band_key(txt) in repeated:
                continue
            kept.append(b)
        p.blocks = kept


def _in_band(block: tuple, p: _Page) -> bool:
    return block[1] < p.height * _HEADER_BAND or block[3] > p.height * _FOOTER_BAND


def _band_key(text: str) -> str:
    """去掉数字与空白后的指纹——页码会变，页眉文字不变。"""
    return re.sub(r"[\d\s\-–—_/·.]+", "", str(text))[:40]


def _is_page_number(text: str) -> bool:
    s = text.strip()
    if len(s) > 12:
        return False
    return bool(re.fullmatch(r"[-–—\s第页共/\\]*\d{1,4}[-–—\s第页共/\\]*", s))


def _blocks_to_text(p: _Page) -> str:
    return "\n".join(str(b[4]).strip() for b in _order_blocks(p.blocks, p.width))


def _order_blocks(blocks: List[tuple], width: float) -> List[tuple]:
    """还原阅读顺序。双栏版式按「先左栏后右栏」，否则按 (y, x)。"""
    if not blocks or width <= 0:
        return blocks
    left = [b for b in blocks if b[2] <= width * 0.52]
    right = [b for b in blocks if b[0] >= width * 0.48]
    is_two_col = (len(left) >= 2 and len(right) >= 2
                  and len(left) + len(right) >= len(blocks) * 0.8)
    if is_two_col:
        return (sorted(left, key=lambda b: (b[1], b[0]))
                + sorted(right, key=lambda b: (b[1], b[0])))
    return sorted(blocks, key=lambda b: (round(b[1] / 6), b[0]))


def _ocr_page(page, name: str, index: int) -> Tuple[str, float, int]:
    try:
        import fitz
        import numpy as np
        pix = page.get_pixmap(dpi=200, colorspace=fitz.csRGB)
        arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width, pix.n)[:, :, :3]
    except Exception as exc:  # noqa: BLE001
        log.warning("PDF 第 %d 页渲染失败（%s）: %s", index + 1, name, exc)
        return "", 0.0, 0
    from app.ingest.ocr import ocr
    return ocr(arr, label=f"{name} 第{index + 1}页")


def _load_pdf_pypdf(path: Path) -> LoadedDoc:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = path.stem
    try:
        if reader.metadata and reader.metadata.title:
            title = str(reader.metadata.title).strip() or title
    except Exception:  # noqa: BLE001
        pass
    pages: List[str] = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            log.warning("PDF 第 %d 页解析失败(%s): %s", i + 1, path.name, exc)
            t = ""
        pages.append(t.strip())
    text = PAGE_BREAK.join(pages)
    text = re.sub(r"(?<=[。！？；])\s*(?=[^\n])", "\n", text)
    meta: Dict[str, object] = {"modality": "pdf", "pages": len(reader.pages)}
    if not text.replace(PAGE_BREAK, "").strip():
        meta["empty_reason"] = "PDF 无文本层（可能是扫描件），且当前环境未装 pymupdf，无法 OCR"
    return LoadedDoc(_clean(_infer_headings_from_plain(text)), title, "pdf", meta)


# ---------------- Office / HTML ----------------

def _load_docx(path: Path) -> LoadedDoc:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    title = path.stem
    blocks: List[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            txt = (p.text or "").strip()
            if not txt:
                continue
            style = (p.style.name or "").lower()
            if "heading" in style or "标题" in style:
                level = 1
                m = re.search(r"(\d+)", style)
                if m:
                    level = min(6, max(1, int(m.group(1))))
                blocks.append(f"{'#' * level} {txt}")
                if level == 1 and title == path.stem:
                    title = txt
            else:
                blocks.append(txt)
        elif tag == "tbl":
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                blocks.append("\n".join(rows))
    return LoadedDoc(_clean("\n\n".join(blocks)), title, "docx", {"modality": "text"})


def _load_pptx(path: Path) -> LoadedDoc:
    from pptx import Presentation

    prs = Presentation(str(path))
    title = path.stem
    blocks: List[str] = []
    for i, slide in enumerate(prs.slides, 1):
        s_title = ""
        try:
            if slide.shapes.title is not None:
                s_title = (slide.shapes.title.text or "").strip()
        except Exception:  # noqa: BLE001
            pass
        if i == 1 and s_title:
            title = s_title
        blocks.append(f"## {s_title or f'第{i}页'}")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                txt = "".join(r.text for r in para.runs).strip()
                if not txt or txt == s_title:
                    continue
                indent = "  " * min(3, para.level or 0)
                blocks.append(f"{indent}- {txt}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                blocks.append(f"讲稿：{note}")
    return LoadedDoc(_clean("\n".join(blocks)), title, "pptx",
                     {"modality": "text", "slides": len(prs.slides)})


def _load_xlsx(path: Path) -> LoadedDoc:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: List[str] = []
    for ws in wb.worksheets:
        blocks.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if not any(cells):
                continue
            blocks.append("| " + " | ".join(cells) + " |")
    wb.close()
    return LoadedDoc(_clean("\n".join(blocks)), path.stem, "xlsx", {"modality": "text"})


def _load_html(path: Path) -> LoadedDoc:
    raw = _read_text(path)
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(p|div|li|tr|h[1-6])>", "\n", raw)
    raw = re.sub(r"(?i)<h([1-6])[^>]*>", lambda m: "\n" + "#" * int(m.group(1)) + " ", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    text = html_lib.unescape(raw)
    title = path.stem
    m = re.search(r"^#\s+(.+)$", text, re.M)
    if m:
        title = m.group(1).strip()
    return LoadedDoc(_clean(text), title, "html", {"modality": "text"})


# ---------------- 图片 ----------------

def _load_image(path: Path) -> LoadedDoc:
    from app.ingest.ocr import available, ocr_path

    if not available():
        raise ValueError("未安装 OCR 依赖（rapidocr-onnxruntime），无法解析图片")
    text, conf, lines = ocr_path(path)
    meta: Dict[str, object] = {
        "modality": "image", "ocr": True,
        "ocr_confidence": conf, "ocr_lines": lines,
    }
    if not text.strip():
        # 不抛异常：图上确实可能没有字。但要留痕，让入库报告把它单列出来。
        meta["empty_reason"] = "图上未识别到文字（若是流程图/架构图，建议补一句图注再用）"
    return LoadedDoc(_clean(_infer_headings_from_plain(text)), path.stem, "image", meta)


# ---------------- 目录扫描 ----------------

def reject_reason(path: Path) -> str | None:
    """支持则返回 None，否则返回人类可读的原因。"""
    ext = path.suffix.lower()
    if ext in SUPPORTED_EXTS:
        return None
    if ext in AUDIO_VIDEO_EXTS:
        return "音视频（本轮未接入，需先转写成文本）"
    if ext in LEGACY_OFFICE_EXTS:
        return f"旧版二进制 Office 格式（{ext}），请另存为 .docx / .xlsx / .pptx"
    if ext in ARCHIVE_EXTS:
        return "压缩包，请解压后入库"
    if not ext:
        return "无扩展名，无法判断格式"
    return f"不支持的扩展名（{ext}）"


def scan_dir(root: Path) -> Tuple[List[Path], List[Tuple[Path, str]]]:
    """扫描待入库文件，返回 (支持的, [(被拒绝的, 原因)])。

    和原来只返回列表的区别：**被拒绝的文件必须能被报出来**。
    原实现按扩展名过滤，非文本文件直接消失，入库报告显示一切正常——
    这是知识库"内容莫名缺失"最常见的来源。
    """
    if root.is_file():
        reason = reject_reason(root)
        return ([] if reason else [root]), ([(root, reason)] if reason else [])

    files: List[Path] = []
    rejected: List[Tuple[Path, str]] = []
    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        if any(part.startswith(".") for part in p.parts):
            continue
        if p.name.startswith("~$"):
            continue
        reason = reject_reason(p)
        if reason:
            rejected.append((p, reason))
        else:
            files.append(p)
    return files, rejected


def walk_files(root: Path) -> List[Path]:
    """兼容旧接口：只返回支持的文件。新代码请用 scan_dir 以拿到被拒绝的清单。"""
    return scan_dir(root)[0]
